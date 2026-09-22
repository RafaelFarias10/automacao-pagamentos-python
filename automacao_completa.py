import re
import xml.etree.ElementTree as ET
from pathlib import Path
from collections import defaultdict
from datetime import datetime

from openpyxl import load_workbook


# ============================================================
# CONFIGURAÇÕES
# ============================================================

PASTA_ENTRADA = Path("entrada")
PASTA_SAIDA = Path("saida")

ARQUIVO_PROTHEUS = PASTA_ENTRADA / "protheus.xls"
ARQUIVO_PLANILHA = PASTA_ENTRADA / "registro_pagamentos.xlsx"

ARQUIVO_XML_LIMPO = PASTA_SAIDA / "protheus_limpo.xml"
ARQUIVO_SAIDA = PASTA_SAIDA / "registro_pagamentos_AUTOMATIZADO.xlsx"

NOME_ABA = "采购付款台账 | Registro de Pagamentos de Compras"

LINHA_INICIAL = 5

NS = {
    "ss": "urn:schemas-microsoft-com:office:spreadsheet"
}


# ============================================================
# FUNÇÕES AUXILIARES
# ============================================================

def limpar_texto(valor):
    if valor is None:
        return ""

    return str(valor).strip()


def normalizar_po(valor):
    """
    Padroniza o PO para comparação.
    Remove espaços e transforma em maiúsculo.
    """
    valor = limpar_texto(valor)
    valor = re.sub(r"\s+", "", valor)

    return valor.upper()


def ler_linha_xml(linha):
    valores = []
    proxima_coluna = 1

    for celula in linha.findall("ss:Cell", NS):

        indice = celula.attrib.get(
            "{urn:schemas-microsoft-com:office:spreadsheet}Index"
        )

        if indice is not None:
            indice = int(indice)

            while proxima_coluna < indice:
                valores.append("")
                proxima_coluna += 1

        dado = celula.find("ss:Data", NS)

        if dado is not None:
            valores.append(dado.text or "")
        else:
            valores.append("")

        proxima_coluna += 1

    return valores


def preencher_se_vazio(celula, valor):
    """
    Só preenche células realmente vazias.
    Nunca substitui informação existente.
    """
    if limpar_texto(valor) == "":
        return False

    if celula.value is None or limpar_texto(celula.value) == "":
        celula.value = valor
        return True

    return False


# ============================================================
# VALIDAÇÃO DOS ARQUIVOS
# ============================================================

print("=" * 55)
print("AUTOMAÇÃO - REGISTRO DE PAGAMENTOS")
print("=" * 55)

PASTA_SAIDA.mkdir(exist_ok=True)

if not ARQUIVO_PROTHEUS.exists():
    raise FileNotFoundError(
        f"\nArquivo não encontrado:\n{ARQUIVO_PROTHEUS}\n"
        "\nColoque a exportação do Protheus na pasta entrada."
    )

if not ARQUIVO_PLANILHA.exists():
    raise FileNotFoundError(
        f"\nArquivo não encontrado:\n{ARQUIVO_PLANILHA}\n"
        "\nColoque a planilha oficial na pasta entrada."
    )


# ============================================================
# ETAPA 1 - LIMPAR XML EXPORTADO PELO PROTHEUS
# ============================================================

print("\n[1/5] Preparando exportação do Protheus...")

# Caracteres que não são permitidos pelo XML 1.0
padrao_invalidos = re.compile(
    rb"[\x00-\x08\x0B\x0C\x0E-\x1F]"
)

bytes_processados = 0

with open(ARQUIVO_PROTHEUS, "rb") as entrada, \
        open(ARQUIVO_XML_LIMPO, "wb") as saida:

    while True:

        bloco = entrada.read(1024 * 1024)

        if not bloco:
            break

        bloco = padrao_invalidos.sub(b"", bloco)

        # Corrige & que não representa entidade XML válida
        bloco = re.sub(
            rb"&(?!amp;|lt;|gt;|quot;|apos;|#[0-9]+;|#x[0-9A-Fa-f]+;)",
            b"&amp;",
            bloco
        )

        saida.write(bloco)

        bytes_processados += len(bloco)

print("Exportação preparada com sucesso.")


# ============================================================
# ETAPA 2 - LER XML
# ============================================================

print("\n[2/5] Lendo dados do Protheus...")

tree = ET.parse(ARQUIVO_XML_LIMPO)
root = tree.getroot()

worksheet = root.find(".//ss:Worksheet", NS)

if worksheet is None:
    raise Exception("Nenhuma planilha encontrada no relatório do Protheus.")

tabela = worksheet.find("ss:Table", NS)

if tabela is None:
    raise Exception("Tabela do relatório do Protheus não encontrada.")

linhas = tabela.findall("ss:Row", NS)

if len(linhas) < 3:
    raise Exception("O relatório do Protheus não contém registros suficientes.")

# Linha 1 = título
# Linha 2 = cabeçalho
cabecalho = ler_linha_xml(linhas[1])

indice_colunas = {
    limpar_texto(nome): indice
    for indice, nome in enumerate(cabecalho)
}

colunas_obrigatorias = [
    "NOME FILIAL",
    "DATA EMISSAO",
    "NOME FORNECEDOR",
    "PRODUTO DESC",
    "PEDIDO",
    "VLR. TOTAL PC"
]

for coluna in colunas_obrigatorias:

    if coluna not in indice_colunas:
        raise Exception(
            f"Coluna obrigatória não encontrada no Protheus: {coluna}"
        )

print(f"Registros encontrados: {len(linhas) - 2}")


# ============================================================
# ETAPA 3 - ORGANIZAR POR PO + REMOVER DUPLICIDADES
# ============================================================

print("\n[3/5] Organizando registros por PO...")

registros_por_po = defaultdict(list)
chaves_processadas = set()

for numero, linha in enumerate(linhas[2:], start=1):

    valores = ler_linha_xml(linha)

    if len(valores) < len(cabecalho):
        valores.extend(
            [""] * (len(cabecalho) - len(valores))
        )

    def pegar(nome_coluna):
        return limpar_texto(
            valores[indice_colunas[nome_coluna]]
        )

    registro = {
        "empresa": pegar("NOME FILIAL"),
        "data_compra": pegar("DATA EMISSAO"),
        "fornecedor": pegar("NOME FORNECEDOR"),
        "descricao": pegar("PRODUTO DESC"),
        "po": pegar("PEDIDO"),
        "valor": pegar("VLR. TOTAL PC"),
    }

    po_normalizado = normalizar_po(registro["po"])

    if po_normalizado == "":
        continue

    # IMPORTANTE:
    # NF não faz parte desta versão.
    chave = (
        registro["empresa"],
        registro["data_compra"],
        registro["fornecedor"],
        registro["descricao"],
        po_normalizado,
        registro["valor"],
    )

    if chave in chaves_processadas:
        continue

    chaves_processadas.add(chave)

    registros_por_po[po_normalizado].append(registro)

    if numero % 10000 == 0:
        print(f"   {numero} registros analisados...")

print("Base organizada.")


# ============================================================
# ETAPA 4 - ABRIR PLANILHA OFICIAL
# ============================================================

print("\n[4/5] Preenchendo planilha oficial...")

wb = load_workbook(ARQUIVO_PLANILHA)

if NOME_ABA not in wb.sheetnames:
    raise Exception(
        f"Aba não encontrada:\n{NOME_ABA}"
    )

ws = wb[NOME_ABA]


# ============================================================
# CONTADORES
# ============================================================

total_pos = 0
automaticos = 0
nao_encontrados = 0
multiplos = 0
celulas_preenchidas = 0


# ============================================================
# PROCURAR CADA PO
# ============================================================

for linha_excel in range(LINHA_INICIAL, ws.max_row + 1):

    # Coluna I = Nº Pedido de Compra (PO)
    po_original = ws.cell(
        row=linha_excel,
        column=9
    ).value

    po = normalizar_po(po_original)

    if po == "":
        continue

    total_pos += 1

    registros = registros_por_po.get(po, [])


    # --------------------------------------------------------
    # NÃO ENCONTRADO
    # --------------------------------------------------------

    if len(registros) == 0:

        nao_encontrados += 1
        print(f"NÃO ENCONTRADO: {po_original!r}")

        continue


    # --------------------------------------------------------
    # MAIS DE UM REGISTRO DIFERENTE
    # --------------------------------------------------------

    if len(registros) > 1:

        multiplos += 1

        print(
            f"MÚLTIPLOS: {po_original!r} "
            f"→ {len(registros)} registros"
        )

        continue


    # --------------------------------------------------------
    # EXATAMENTE UM REGISTRO
    # --------------------------------------------------------

    registro = registros[0]

    automaticos += 1

    # B - Nome da Empresa
    if preencher_se_vazio(
        ws.cell(linha_excel, 2),
        registro["empresa"]
    ):
        celulas_preenchidas += 1


    # C - Data de Compra
    celula_data = ws.cell(linha_excel, 3)

    if (
        celula_data.value is None
        or limpar_texto(celula_data.value) == ""
    ):

        data = registro["data_compra"]

        if data:

            try:
                data_excel = datetime.strptime(
                    data[:10],
                    "%Y-%m-%d"
                )

                celula_data.value = data_excel
                celula_data.number_format = "dd/mm/yyyy"

                celulas_preenchidas += 1

            except ValueError:
                pass


    # F - Nome do Fornecedor
    if preencher_se_vazio(
        ws.cell(linha_excel, 6),
        registro["fornecedor"]
    ):
        celulas_preenchidas += 1


    # G - Descrição da Compra
    if preencher_se_vazio(
        ws.cell(linha_excel, 7),
        registro["descricao"]
    ):
        celulas_preenchidas += 1


    # I = PO
    # Não alteramos porque ele já veio da planilha oficial.


    # K = Nota Fiscal
    # NÃO PREENCHER.
    # Será tratada futuramente pelo Banco de Conhecimento.


    # N - Valor Faturado sem imposto
    celula_valor = ws.cell(linha_excel, 14)

    if (
        celula_valor.value is None
        or limpar_texto(celula_valor.value) == ""
    ):

        valor = registro["valor"]

        if valor:

            try:
                celula_valor.value = float(valor)
                celula_valor.number_format = '#,##0.00'

                celulas_preenchidas += 1

            except ValueError:
                pass


# ============================================================
# ETAPA 5 - SALVAR
# ============================================================

print("\n[5/5] Salvando resultado...")

try:

    wb.save(ARQUIVO_SAIDA)

except PermissionError:

    print("\nERRO:")
    print("O arquivo de saída provavelmente está aberto no Excel.")
    print("Feche o arquivo e execute novamente.")

    raise


# ============================================================
# RESULTADO
# ============================================================

print("\n" + "=" * 55)
print("AUTOMAÇÃO FINALIZADA")
print("=" * 55)

print(f"POs analisados:                 {total_pos}")
print(f"POs automáticos:                {automaticos}")
print(f"POs não encontrados:            {nao_encontrados}")
print(f"POs com múltiplos registros:    {multiplos}")
print(f"Células novas preenchidas:      {celulas_preenchidas}")

print("-" * 55)

print("Arquivo criado:")
print(ARQUIVO_SAIDA)

print("=" * 55)
